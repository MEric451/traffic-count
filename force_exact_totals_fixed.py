import openpyxl
from openpyxl import load_workbook
import random

def force_exact_totals(input_file, output_file, totals):
    """
    Force EXACT totals to match the hardcoded values (12-hour version)
    """
    wb = load_workbook(input_file)
    
    print("FORCING EXACT TOTALS TO MATCH HARDCODED VALUES (12-HOUR)")
    print("="*50)
    print(f"All sheets in workbook: {wb.sheetnames}")
    
    # Only process 12-hour sheets (7AM-7PM)
    valid_sheets = ['7-8AM', '8-9AM', '9-10AM', '10-11AM', '11-12PM', '12-1PM', 
                    '1-2PM', '2-3PM', '3-4PM', '4-5PM', '5-6PM', '6-7PM']
    hourly_sheets = [name for name in wb.sheetnames if name in valid_sheets]
    
    print(f"Processing these {len(hourly_sheets)} sheets: {hourly_sheets}")
    
    # Find direction rows from reference sheet
    reference_sheet = wb['7-8AM']
    direction_rows = {}
    
    for row in range(1, 100):
        cell_a = reference_sheet.cell(row=row, column=1).value
        if cell_a:
            cell_a_clean = ' '.join(str(cell_a).lower().split())
            for direction in totals.keys():
                direction_clean = ' '.join(direction.lower().split())
                if direction_clean in cell_a_clean:
                    cell_b = reference_sheet.cell(row=row, column=2).value
                    if not (isinstance(cell_b, str) and cell_b.startswith('=')):
                        if direction not in direction_rows:
                            direction_rows[direction] = []
                        direction_rows[direction].append(row)
    
    print("\nFound direction rows:")
    for direction, rows in direction_rows.items():
        print(f"  {direction}: {rows}")
    
    # Process each direction and vehicle class
    for direction, direction_totals in totals.items():
        print(f"\nProcessing {direction}...")
        
        if direction not in direction_rows:
            print(f"  WARNING: {direction} not found in Excel file!")
            continue
            
        target_rows = direction_rows[direction]
        
        for col_idx, exact_total in enumerate(direction_totals):
            print(f"  Class {col_idx+1}: Forcing EXACT {exact_total} vehicles")
            
            if exact_total == 0:
                # Set all hours to 0
                for sheet_name in hourly_sheets:
                    sheet = wb[sheet_name]
                    for row_num in target_rows:
                        sheet.cell(row=row_num, column=col_idx + 2).value = 0
                continue
            
            # Create simple hourly distribution without random variations
            num_hours = len(hourly_sheets)
            base_per_hour = exact_total // num_hours
            remainder = exact_total % num_hours
            
            hourly_distribution = {}
            total_allocated = 0
            
            # Distribute evenly with remainder
            for i, sheet_name in enumerate(hourly_sheets):
                if i < remainder:
                    hourly_value = base_per_hour + 1
                else:
                    hourly_value = base_per_hour
                
                hourly_distribution[sheet_name] = hourly_value
                total_allocated += hourly_value
            
            # CRITICAL: Verify exact total
            if total_allocated != exact_total:
                print(f"    ERROR: Allocated {total_allocated} but need {exact_total}")
                # Force correction
                difference = exact_total - total_allocated
                first_sheet = hourly_sheets[0]
                hourly_distribution[first_sheet] += difference
                print(f"    FIXED: Added {difference} to {first_sheet}")
            
            # Apply to sheets - distribute across rows
            for sheet_name in hourly_sheets:
                sheet = wb[sheet_name]
                hourly_value = hourly_distribution[sheet_name]
                
                if len(target_rows) > 0 and hourly_value > 0:
                    # Simple distribution across rows
                    base_per_row = hourly_value // len(target_rows)
                    row_remainder = hourly_value % len(target_rows)
                    
                    remaining = hourly_value
                    for i, row_num in enumerate(target_rows):
                        if i < row_remainder:
                            value = base_per_row + 1
                        else:
                            value = base_per_row
                        
                        # Ensure we don't exceed remaining
                        if i == len(target_rows) - 1:
                            value = remaining  # Last row gets whatever is left
                        
                        sheet.cell(row=row_num, column=col_idx + 2).value = value
                        remaining -= value
                else:
                    # No vehicles for this hour
                    for row_num in target_rows:
                        sheet.cell(row=row_num, column=col_idx + 2).value = 0
    
    wb.save(output_file)
    print(f"\nSaved to: {output_file}")
    
    verify_exact_totals(wb, totals, hourly_sheets, direction_rows)

def verify_exact_totals(wb, totals, hourly_sheets, direction_rows):
    """Verify the totals are EXACTLY what we wanted"""
    print("\nFINAL VERIFICATION - EXACT MATCH CHECK:")
    print("="*50)
    
    all_perfect = True
    
    for direction, expected_totals in totals.items():
        if direction not in direction_rows:
            print(f"SKIPPED: {direction} (not found in Excel)")
            continue
            
        target_rows = direction_rows[direction]
        print(f"\n{direction}:")
        
        direction_actual = 0
        direction_expected = sum(expected_totals)
        
        for col_idx, expected_class_total in enumerate(expected_totals):
            class_actual = 0
            
            # Sum across all hours and rows for this vehicle class
            for sheet_name in hourly_sheets:
                sheet = wb[sheet_name]
                for row_num in target_rows:
                    val = sheet.cell(row=row_num, column=col_idx + 2).value or 0
                    class_actual += val
            
            direction_actual += class_actual
            
            if class_actual == expected_class_total:
                status = "✓ PERFECT"
            else:
                status = f"✗ ERROR (got {class_actual}, diff: {class_actual - expected_class_total})"
                all_perfect = False
            
            print(f"  Class {col_idx+1}: Expected {expected_class_total} → Got {class_actual} [{status}]")
        
        total_status = "✓ PERFECT MATCH" if direction_actual == direction_expected else f"✗ TOTAL ERROR"
        print(f"  DIRECTION TOTAL: Expected {direction_expected} → Got {direction_actual} [{total_status}]")
    
    if all_perfect:
        print(f"\n🎉 SUCCESS: ALL TOTALS ARE EXACTLY AS SPECIFIED!")
    else:
        print(f"\n❌ ERROR: Some totals don't match. Check the output above.")

if __name__ == "__main__":
    # Your EXACT totals - these will be forced exactly
    totals = {
        "Mombasa Bound Lane 1": [1081, 820, 429, 135, 230, 36, 13, 13, 133, 36, 46, 3],
        "Mombasa Bound Lane 2": [447, 996, 658, 260, 235, 97, 31, 25, 600, 339, 1055, 3],
        "Mombasa Bound Lane 3": [97, 1419, 1038, 391, 277, 35, 17, 33, 492, 164, 635, 3]
    }

    input_file = "Athi River Site, Day 3, Wednesday, Nairobi - Mombasa (A8) Road Traffic Volumes Hourly Tally.xlsx"
    output_file = "Mogogosiek Site, DAY 6 Saturday Litein - Silibwet - Bomet (B124) Road Site Traffic Volumes Hourly Traffic Tally.xlsx"
    
    force_exact_totals(input_file, output_file, totals)